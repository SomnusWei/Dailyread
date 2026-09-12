#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
医学知识库文本读取工具（共享模块）

支持格式：txt / pdf / docx / doc
- txt: 自动编码检测（UTF-8 / GB18030）
- pdf: 文本层优先 + 扫描页 OCR 兜底 + 内嵌图片 OCR
- docx: python-docx 读取段落与表格
- doc: 通过 Word COM 转 docx 后读取（需本机安装 Word）

用法（其他脚本）：
  from text_utils import read_text_lines, read_text, detect_encoding
  from text_utils import extract_file_text, extract_pdf_text, extract_docx_text
"""

import os

# 编码探测顺序：严格 UTF-8 → 严格 GB18030 → 容错 GB18030
_CANDIDATE_ENCODINGS = ("utf-8", "gb18030")


def detect_encoding(path):
    """探测文件编码，返回编码名；无法确定时返回 'gb18030'（容错兜底）"""
    try:
        with open(path, "rb") as f:
            data = f.read()
    except OSError:
        return "utf-8"
    for enc in _CANDIDATE_ENCODINGS:
        try:
            data.decode(enc)
            return enc
        except (UnicodeDecodeError, ValueError):
            continue
    return "gb18030"  # 含损坏字节，交给调用方容错


def read_text(path, errors="strict"):
    """读取整个文件文本，自动处理 UTF-8 / GB18030 编码。

    errors='strict'（默认）：损坏字节会抛 UnicodeDecodeError
    errors='replace'      ：损坏字节替换为 U+FFFD，保证不中断
    """
    enc = detect_encoding(path)
    if errors != "strict":
        enc = enc if enc in _CANDIDATE_ENCODINGS else "gb18030"
        with open(path, "r", encoding=enc, errors=errors) as f:
            return f.read()
    try:
        with open(path, "r", encoding=enc) as f:
            return f.read()
    except UnicodeDecodeError:
        with open(path, "r", encoding="gb18030", errors="replace") as f:
            return f.read()


def read_text_lines(path):
    """读取文件为行列表（保留换行符），自动编码检测 + 容错。"""
    return read_text(path, errors="replace").splitlines(keepends=True)


def clean_ancient_markup(text):
    """清理古籍文本标记：<篇名>/<目录>/属性：/\r"""
    text = text.replace("<篇名>", "【篇】").replace("<目录>", "【目录】")
    text = text.replace("属性：", "")
    return text.replace("\r", "")


# ===================== PDF 文本提取 =====================

def _open_fitz():
    """兼容导入 PyMuPDF：新版推荐 pymupdf，旧版用 fitz。"""
    try:
        import pymupdf as fitz_mod  # type: ignore
        return fitz_mod
    except ImportError:
        import fitz  # type: ignore  # noqa: F401
        return fitz


def _init_ocr():
    """惰性初始化 RapidOCR 引擎。成功返回引擎对象，失败返回 None。"""
    try:
        from rapidocr import RapidOCR  # type: ignore
        return RapidOCR()
    except Exception:
        return None


def _ocr_image(ocr_engine, image_bytes):
    """对单张 PNG/JPG 字节做 OCR，返回识别文本（按行拼接）。"""
    import numpy as np
    from io import BytesIO
    from PIL import Image
    img = Image.open(BytesIO(image_bytes)).convert("RGB")
    arr = np.array(img)
    result = ocr_engine(arr)
    # rapidocr >= 3.0 返回 RapidOCROutput，含 .txts
    txts = getattr(result, "txts", None)
    if txts is None and isinstance(result, (list, tuple)) and result:
        txts = [r[1] for r in result if r]
    if not txts:
        return ""
    return "\n".join(t for t in txts if t)


def extract_pdf_text(pdf_path, dpi=200, min_text_chars=30, use_ocr=True,
                     ocr_engine=None, extract_images=True):
    """从 PDF 提取文本，优先用文本层，扫描页用 OCR 兜底，可选提取内嵌图片 OCR。

    Args:
        pdf_path: PDF 文件路径
        dpi: 渲染页面供 OCR 的分辨率（默认 200）
        min_text_chars: 单页文本层字符数低于此值则判定为扫描页，触发整页 OCR
        use_ocr: 是否启用 OCR 兜底（False 时只提取文本层）
        ocr_engine: 复用已初始化的 RapidOCR 引擎；为 None 时按需初始化
        extract_images: 是否对 PDF 内嵌图片（图表/解剖图等）做 OCR（默认 True）

    Returns:
        (text, stats) 元组：
            text: 合并后的全文本字符串
            stats: dict，含 pages / text_pages / ocr_pages / images_ocr'd / source
    """
    fitz = _open_fitz()
    doc = fitz.open(pdf_path)
    pages_text = []
    text_pages = 0
    ocr_pages = 0
    images_ocr = 0

    for page_idx, page in enumerate(doc):
        raw = page.get_text("text") or ""
        page_parts = []

        # 1. 文本层
        if len(raw.strip()) >= min_text_chars:
            page_parts.append(raw)
            text_pages += 1
        else:
            # 2. 扫描页：整页渲染 OCR
            if use_ocr:
                if ocr_engine is None:
                    ocr_engine = _init_ocr()
                if ocr_engine is not None:
                    pix = page.get_pixmap(dpi=dpi)
                    png_bytes = pix.tobytes("png")
                    ocr_text = _ocr_image(ocr_engine, png_bytes)
                    page_parts.append(ocr_text if ocr_text else raw)
                    ocr_pages += 1
                else:
                    page_parts.append(raw)
            else:
                page_parts.append(raw)

        # 3. 内嵌图片 OCR（教材中的图表、解剖图、表格截图）
        if extract_images and use_ocr:
            if ocr_engine is None:
                ocr_engine = _init_ocr()
            if ocr_engine is not None:
                image_list = page.get_images(full=True)
                for img_index, img in enumerate(image_list):
                    xref = img[0]
                    try:
                        base_image = doc.extract_image(xref)
                        img_bytes = base_image["image"]
                        # 跳过过小的图片（可能是装饰性图标）
                        if len(img_bytes) < 2048:
                            continue
                        ocr_result = _ocr_image(ocr_engine, img_bytes)
                        if ocr_result.strip():
                            page_parts.append(f"[图片{img_index + 1}]\n{ocr_result}")
                            images_ocr += 1
                    except Exception:
                        continue

        pages_text.append("\n".join(p for p in page_parts if p))

    doc.close()
    full_text = "\n\n".join(pages_text)
    if ocr_pages > 0 or images_ocr > 0:
        source = "ocr"
    elif text_pages > 0:
        source = "text"
    else:
        source = "empty"
    stats = {
        "pages": len(pages_text),
        "text_pages": text_pages,
        "ocr_pages": ocr_pages,
        "images_ocr": images_ocr,
        "source": source,
    }
    return full_text, stats


# ===================== DOCX 文本提取 =====================

def extract_docx_text(docx_path):
    """从 .docx 提取文本（段落 + 表格）。

    Returns:
        (text, stats) 元组
    """
    from docx import Document
    doc = Document(docx_path)
    parts = []

    # 段落
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # 表格
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    full_text = "\n".join(parts)
    stats = {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables),
             "source": "docx"}
    return full_text, stats


# ===================== DOC 文本提取（旧版 Word） =====================

def extract_doc_text(doc_path):
    """从 .doc 提取文本：通过 Word COM 转存为 docx 后读取。

    要求本机安装 Microsoft Word。失败时返回空文本和错误信息。

    Returns:
        (text, stats) 元组
    """
    import tempfile
    try:
        import win32com.client as win32
    except ImportError:
        return "", {"source": "doc", "error": "pywin32 未安装，无法处理 .doc"}

    word = None
    tmp_docx = None
    try:
        word = win32.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(os.path.abspath(doc_path))
        tmp_dir = tempfile.mkdtemp()
        tmp_docx = os.path.join(tmp_dir, "_converted.docx")
        # 16 = wdFormatDocumentDefault (.docx)
        doc.SaveAs(os.path.abspath(tmp_docx), FileFormat=16)
        doc.Close(False)
        word.Quit()
        word = None
        return extract_docx_text(tmp_docx)
    except Exception as e:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        return "", {"source": "doc", "error": str(e)}
    finally:
        if tmp_docx and os.path.exists(tmp_docx):
            try:
                os.remove(tmp_docx)
            except Exception:
                pass


# ===================== TXT 文本提取 =====================

def extract_txt_text(txt_path):
    """从 .txt 提取文本（自动编码检测）。

    Returns:
        (text, stats) 元组
    """
    text = read_text(txt_path, errors="replace")
    stats = {"source": "txt", "encoding": detect_encoding(txt_path)}
    return text, stats


# ===================== 统一入口 =====================

# 支持的扩展名
SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx", ".doc"}


def extract_file_text(file_path, **kwargs):
    """统一文件文本提取入口，根据扩展名自动分发。

    Args:
        file_path: 文件路径
        **kwargs: 透传给对应提取函数（如 dpi, use_ocr, extract_images）

    Returns:
        (text, stats) 元组；不支持的格式返回 ("", {"error": "unsupported"})
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".txt":
        return extract_txt_text(file_path)
    elif ext == ".pdf":
        return extract_pdf_text(file_path, **kwargs)
    elif ext == ".docx":
        return extract_docx_text(file_path)
    elif ext == ".doc":
        return extract_doc_text(file_path)
    else:
        return "", {"error": f"unsupported extension: {ext}",
                    "supported": sorted(SUPPORTED_EXTENSIONS)}


if __name__ == "__main__":
    # 自检：对目录下所有 txt 做编码探测
    import sys
    import collections
    directory = sys.argv[1] if len(sys.argv) > 1 else os.path.abspath(
        os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "中医教材"))
    stats = collections.Counter()
    for fname in sorted(os.listdir(directory)):
        if not fname.lower().endswith(".txt"):
            continue
        enc = detect_encoding(os.path.join(directory, fname))
        stats[enc] += 1
    for enc, n in stats.most_common():
        print(f"{enc}: {n} 个文件")
